#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_contact_book 的回歸測試。

重點是第三項：證明 schema 檢查能攔下「把元素 append 到 w:tblPr 尾端」
這個 bug——它就是 2026-09-22 產出無法開啟的同一類成因，而且檔案大小
與 ZIP 結構都正常，光看這兩者驗不出來。

執行：python3 test_build.py
"""

import json
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).parent))
import build_contact_book as B  # noqa: E402


def _sample():
    sample = Path(__file__).resolve().parent.parent / "references" / "entries.example.json"
    return json.loads(sample.read_text("utf-8"))


def test_full_sheet_passes():
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "a.docx"
        B.build(_sample(), out)
        problems = B.verify(out)
        assert problems == [], problems
        assert zipfile.ZipFile(out).testzip() is None
    print("✓ 8 格完整表格：ZIP、schema、尺寸皆通過")


def test_partial_sheet_keeps_eight_cells():
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "b.docx"
        B.build(_sample()[:3], out)
        assert B.verify(out) == []
        t = Document(out).tables[0]
        assert len(t.rows) * len(t.columns) == 8
        assert t.cell(0, 0).text.startswith("S01今日表現：")
        assert t.cell(3, 1).text.strip() == ""      # 未填格留白
    print("✓ 僅 3 筆輸入：仍輸出 8 格，未填格留白")


def test_schema_checker_catches_appended_element():
    """把 tblCellMar append 到尾端（原始 bug），檢查器必須抓到。"""
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "c.docx"
        B.build(_sample(), out)

        doc = Document(out)
        tbl_pr = doc.tables[0]._tbl.tblPr
        stray = tbl_pr.makeelement(qn("w:tblCellMar"), {})   # 已存在 → 重複 + 順序錯
        tbl_pr.append(stray)                                  # 排在 tblLook 之後
        doc.save(out)

        assert zipfile.ZipFile(out).testzip() is None, "ZIP 仍完整——正是難以察覺之處"
        problems = B.verify(out)
        assert any("順序" in p or "重複" in p for p in problems), problems
    print("✓ 回歸防護：append 到 w:tblPr 尾端會被 schema 檢查攔下")
    print(f"    攔下訊息：{problems[0]}")


def test_independence_blocks_cross_reference():
    """重現 2026-09-22 的實際問題：「同樣不太喜歡拔草工作」。

    「同樣」指涉的是另一格學生的敘述；該生家長只會收到自己孩子那一格，
    讀到時無從得知在跟誰「同樣」。這類敘述必須在產出前被攔下。
    """
    errors, _ = B.check_independence([
        {"label": "S05", "text": "今日課程為柚子娃娃製作。同樣不太喜歡拔草工作，但態度認真。"},
    ])
    assert len(errors) == 1 and "同樣" in errors[0], errors
    print("✓ 逐格獨立性：跨生指涉「同樣」被攔下")


def test_independence_flags_other_cross_ref_forms():
    cases = ["跟同學一樣喜歡畫畫", "其他同學都完成了", "相較之下更專注", "比同學更快完成"]
    for text in cases:
        errors, _ = B.check_independence([{"label": "S01", "text": text}])
        assert errors, f"未攔下：{text}"
    print(f"✓ 逐格獨立性：另外 {len(cases)} 種跨生句型皆被攔下")


def test_independence_allows_clean_text():
    """正常敘述不得誤判——提及「同學協助」是描述支持方式，非跨格指涉。"""
    errors, hints = B.check_independence(_sample())
    assert errors == [], errors
    errors2, _ = B.check_independence([
        {"label": "S06", "text": "生活課練習擦桌子，在老師與同學協助下順利完成整理工作。"},
    ])
    assert errors2 == [], errors2
    print("✓ 逐格獨立性：範例檔與「同學協助」等正常敘述無誤判")


def test_independence_hints_are_advisory():
    """比較語只提示、不阻擋，因為可能只是描述該生自身偏好。"""
    errors, hints = B.check_independence([
        {"label": "S02", "text": "園藝課表現最為積極。"},
    ])
    assert errors == []
    assert len(hints) == 1 and "最為" in hints[0], hints
    print("✓ 逐格獨立性：比較語僅提示，且不重複計數")


def test_cell_width_is_8_5cm():
    from docx.shared import Cm
    assert B.LAYOUT["cell_w_cm"] == 8.5
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "w.docx"
        B.build(_sample(), out)
        assert B.verify(out) == []
        t = Document(out).tables[0]
        assert abs(t.columns[0].width - Cm(8.5)) <= Cm(0.01)
    print(f"✓ 版面：每格 8.5cm × 5.0cm，單格容量 {B.MAX_CHARS} 字")


def test_overlength_warns_but_still_builds():
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "e.docx"
        warnings = B.build([{"label": "S01", "text": "字" * 300}], out)
        assert warnings and "超出" in warnings[0]
        assert B.verify(out) == []
    print("✓ 超長內容：發出篇幅警告，檔案本身仍有效")


if __name__ == "__main__":
    for fn in (
        test_full_sheet_passes,
        test_partial_sheet_keeps_eight_cells,
        test_schema_checker_catches_appended_element,
        test_independence_blocks_cross_reference,
        test_independence_flags_other_cross_ref_forms,
        test_independence_allows_clean_text,
        test_independence_hints_are_advisory,
        test_cell_width_is_8_5cm,
        test_overlength_warns_but_still_builds,
    ):
        fn()
    print("\n全部通過。")
